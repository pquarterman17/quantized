"""Word (.docx) / PowerPoint (.pptx) report writers (ORIGIN_GAP_PLAN #37).

Split out of ``io.report_export`` (PRIMARY_SOFTWARE_AUDIT_PLAN P3.6) when the
figure path grew real embedding. Both writers walk the #36 schema exactly as
before; what changed is the figure block (see ``io.report_figures`` for the
decision rules):

* A block whose ``spec`` the report route rendered embeds the app's own
  figure export as a **300-DPI PNG**, at the figure's natural (requested)
  size, scaled down -- aspect kept -- only to fit the page / slide.
* When that spec requested a **vector** figure, the SVG render is embedded
  too, the way Office itself stores an SVG picture: the PNG stays the
  picture's ``a:blip`` (what Office 2013 and older, and other readers, show)
  and the SVG is attached through the Office 2016+ ``asvg:svgBlip``
  extension on that blip, which Word/PowerPoint 2016+ and Microsoft 365 draw
  instead. If the vector failed to render while the raster succeeded, the
  caption says so -- never a silent downgrade.
* Every embedded picture gets alt text (``descr`` = caption or name) and a
  ``title`` (the figure name) for screen readers.
* A block that cannot be embedded becomes a paragraph / text box naming the
  figure and the specific reason, and a warning is appended to ``warnings``.

``python-docx`` / ``python-pptx`` (MIT) are optional (``quantized[office]``);
the imports are guarded and a missing library raises ``ReportExportError``.
Pure ``io`` layer -- no web-stack imports.
"""

from __future__ import annotations

import io as _io
from collections.abc import Mapping
from typing import Any

from quantized.heavy_import import heavy_imports
from quantized.io.report_blocks import params_rows, table_rows
from quantized.io.report_figures import (
    FigureKey,
    FigurePlan,
    RenderedFigure,
    figure_alt_text,
    figure_warning,
    fit_size,
    placeholder_text,
    plan_figure,
    png_size_in,
)

__all__ = ["OfficeLibraryMissing", "to_docx", "to_pptx"]

_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
# The a:ext uri Office writes for an SVG picture's svgBlip extension.
_SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
_EMU_PER_IN = 914400
_FALLBACK_SIZE_IN = (6.0, 4.5)  # only if a render's PNG header is unreadable


class OfficeLibraryMissing(ImportError):
    """python-docx / python-pptx is not installed (the ``office`` extra)."""


def _svg_ext_xml(rid: str) -> str:
    return (f'<a:ext xmlns:a="{_A_NS}" uri="{_SVG_EXT_URI}">'
            f'<asvg:svgBlip xmlns:asvg="{_ASVG_NS}" xmlns:r="{_R_NS}" r:embed="{rid}"/>'
            "</a:ext>")


def _append_svg_ext(blip: Any, rid: str, parse_xml: Any) -> None:
    """Attach the svgBlip extension to ``blip`` (inside its a:extLst)."""
    ext_lst = blip.find(f"{{{_A_NS}}}extLst")
    if ext_lst is None:
        ext_lst = parse_xml(f'<a:extLst xmlns:a="{_A_NS}"/>')
        blip.append(ext_lst)
    ext_lst.append(parse_xml(_svg_ext_xml(rid)))


def _set_alt(c_nv_pr: Any, block: Mapping[str, Any]) -> None:
    c_nv_pr.set("descr", figure_alt_text(block))
    c_nv_pr.set("title", str(block.get("name", "")))


def _caption_text(block: Mapping[str, Any], note: str | None) -> str:
    text = str(block.get("caption") or block.get("name") or "")
    return f"{text} ({note})" if note else text


def _plan(
    block: Mapping[str, Any],
    key: FigureKey,
    figures: Mapping[FigureKey, RenderedFigure] | None,
    warnings: list[str] | None,
) -> FigurePlan:
    plan = plan_figure(block, (figures or {}).get(key), "office")
    if plan.kind == "render" and plan.png is None:
        plan = FigurePlan("placeholder", reason="no raster (PNG) was rendered for Office")
    if warnings is not None:
        if plan.kind == "placeholder":
            warnings.append(figure_warning(block, key[0], f"not embedded: {plan.reason}"))
        elif plan.note:
            warnings.append(figure_warning(block, key[0], plan.note))
    return plan


def _natural_size(plan: FigurePlan) -> tuple[float, float]:
    assert plan.png is not None
    return png_size_in(plan.png, plan.dpi) or _FALLBACK_SIZE_IN


# ── Word ──────────────────────────────────────────────────────────────────
def _docx_attach_svg(doc: Any, shape: Any, svg: bytes) -> None:
    with heavy_imports("docx.opc.constants", "docx.opc.part", "docx.oxml.parser"):
        from docx.opc.constants import RELATIONSHIP_TYPE
        from docx.opc.part import Part
        from docx.oxml.parser import parse_xml

    part = doc.part
    partname = part.package.next_partname("/word/media/image%d.svg")
    rid = part.relate_to(Part(partname, "image/svg+xml", svg, part.package),
                         RELATIONSHIP_TYPE.IMAGE)
    _append_svg_ext(shape._inline.graphic.graphicData.pic.blipFill.blip, rid, parse_xml)


def _docx_figure(doc: Any, block: Mapping[str, Any], plan: FigurePlan, inches: Any) -> None:
    if plan.kind == "image":  # a caller-supplied raster: pre-P3.6 layout kept
        assert plan.raster is not None
        shape = doc.add_picture(_io.BytesIO(plan.raster), width=inches(6))
        _set_alt(shape._inline.docPr, block)
        if block.get("caption"):
            doc.add_paragraph().add_run(str(block["caption"])).italic = True
        return
    if plan.kind == "placeholder":
        doc.add_paragraph(placeholder_text(block, plan.reason or ""))
        return
    assert plan.png is not None
    sec = doc.sections[-1]
    max_w = (sec.page_width - sec.left_margin - sec.right_margin) / _EMU_PER_IN
    # leave ~0.8 in of the text block for the caption line
    max_h = (sec.page_height - sec.top_margin - sec.bottom_margin) / _EMU_PER_IN - 0.8
    w, h = fit_size(*_natural_size(plan), max_w, max_h)
    shape = doc.add_picture(_io.BytesIO(plan.png), width=inches(w), height=inches(h))
    _set_alt(shape._inline.docPr, block)
    if plan.svg is not None:
        _docx_attach_svg(doc, shape, plan.svg)
    caption = _caption_text(block, plan.note)
    if caption:
        doc.add_paragraph().add_run(caption).italic = True


def to_docx(
    report: Mapping[str, Any],
    *,
    figures: Mapping[FigureKey, RenderedFigure] | None = None,
    warnings: list[str] | None = None,
) -> bytes:
    """The report as a Word document (``figures`` keyed ``(section, block)``)."""
    try:
        with heavy_imports("docx", "docx.shared"):
            from docx import Document  # python-docx (MIT)
            from docx.shared import Inches
    except ImportError as exc:  # pragma: no cover - exercised only without the dep
        raise OfficeLibraryMissing(
            "Word export needs 'python-docx' (pip install quantized[office])"
        ) from exc

    doc = Document()
    doc.add_heading(str(report.get("title", "Report")), level=0)
    for si, sec in enumerate(report.get("sections", [])):
        doc.add_heading(str(sec.get("title", "")), level=1)
        for bi, block in enumerate(sec.get("blocks", [])):
            btype = block.get("type")
            if btype == "text":
                doc.add_paragraph(block["text"])
            elif btype in ("params", "table"):
                header, rows = (params_rows if btype == "params" else table_rows)(block)
                cap = block.get("caption")
                if cap:
                    doc.add_paragraph().add_run(str(cap)).italic = True
                t = doc.add_table(rows=1, cols=len(header))
                t.style = "Light Grid Accent 1"
                for j, h in enumerate(header):
                    t.rows[0].cells[j].text = h
                for row in rows:
                    cells = t.add_row().cells
                    for j, c in enumerate(row):
                        cells[j].text = str(c)
            elif btype == "figure":
                _docx_figure(doc, block, _plan(block, (si, bi), figures, warnings), Inches)
    buf = _io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── PowerPoint ────────────────────────────────────────────────────────────
_SLIDE_TOP_IN = 1.3  # first content row below the section-title box
_SLIDE_BOTTOM_PAD_IN = 0.3
_CAPTION_H_IN = 0.5


def _pptx_attach_svg(slide: Any, picture: Any, svg: bytes) -> None:
    with heavy_imports("pptx.opc.constants", "pptx.opc.package", "pptx.oxml"):
        from pptx.opc.constants import RELATIONSHIP_TYPE
        from pptx.opc.package import Part
        from pptx.oxml import parse_xml

    part = slide.part
    partname = part.package.next_partname("/ppt/media/image%d.svg")
    rid = part.relate_to(Part(partname, "image/svg+xml", part.package, svg),
                         RELATIONSHIP_TYPE.IMAGE)
    _append_svg_ext(picture._element.blipFill.blip, rid, parse_xml)


class _Deck:
    """The presentation being written plus the current slide's fill cursor."""

    def __init__(self, prs: Any, inches: Any, pt: Any) -> None:
        self.prs, self.inches, self.pt = prs, inches, pt
        self.blank = prs.slide_layouts[6]
        self.slide: Any = None
        self.top = _SLIDE_TOP_IN
        self.width_in = prs.slide_width / _EMU_PER_IN
        self.height_in = prs.slide_height / _EMU_PER_IN

    def new_slide(self, title: str) -> None:
        inches = self.inches
        self.slide = self.prs.slides.add_slide(self.blank)
        box = self.slide.shapes.add_textbox(inches(0.5), inches(0.3), inches(9), inches(0.8))
        box.text_frame.text = title
        box.text_frame.paragraphs[0].font.size = self.pt(28)
        self.top = _SLIDE_TOP_IN

    def textbox(self, text: str, height: float, advance: float) -> Any:
        inches = self.inches
        tb = self.slide.shapes.add_textbox(inches(0.5), inches(self.top), inches(9),
                                           inches(height))
        tb.text_frame.text = text
        self.top += advance
        return tb

    def figure(self, block: Mapping[str, Any], plan: FigurePlan, section_title: str) -> None:
        inches = self.inches
        if plan.kind == "image":  # a caller-supplied raster: pre-P3.6 layout kept
            assert plan.raster is not None
            pic = self.slide.shapes.add_picture(
                _io.BytesIO(plan.raster), inches(0.5), inches(self.top), width=inches(6)
            )
            _set_alt(pic._element.nvPicPr.cNvPr, block)
            self.top += 4.0
            return
        if plan.kind == "placeholder":
            self.textbox(placeholder_text(block, plan.reason or ""), 0.5, 0.6)
            return
        assert plan.png is not None
        max_w = self.width_in - 1.0
        room = self.height_in - _SLIDE_BOTTOM_PAD_IN - _CAPTION_H_IN
        w, h = fit_size(*_natural_size(plan), max_w, room - _SLIDE_TOP_IN)
        if self.top + h > room and self.top > _SLIDE_TOP_IN:
            self.new_slide(f"{section_title} (cont.)")  # never overflow the slide
        left = 0.5 + (max_w - w) / 2
        pic = self.slide.shapes.add_picture(
            _io.BytesIO(plan.png), inches(left), inches(self.top),
            width=inches(w), height=inches(h),
        )
        _set_alt(pic._element.nvPicPr.cNvPr, block)
        if plan.svg is not None:
            _pptx_attach_svg(self.slide, pic, plan.svg)
        self.top += h
        caption = _caption_text(block, plan.note)
        if caption:
            tb = self.textbox(caption, 0.4, _CAPTION_H_IN)
            para = tb.text_frame.paragraphs[0]
            para.font.size, para.font.italic = self.pt(12), True
            tb.text_frame.word_wrap = True


def to_pptx(
    report: Mapping[str, Any],
    *,
    figures: Mapping[FigureKey, RenderedFigure] | None = None,
    warnings: list[str] | None = None,
) -> bytes:
    """The report as a slide deck: a title slide, then one slide per section
    (a figure that no longer fits continues on a "(cont.)" slide)."""
    try:
        with heavy_imports("pptx", "pptx.util"):
            from pptx import Presentation  # python-pptx (MIT)
            from pptx.util import Inches, Pt
    except ImportError as exc:  # pragma: no cover - exercised only without the dep
        raise OfficeLibraryMissing(
            "PowerPoint export needs 'python-pptx' (pip install quantized[office])"
        ) from exc

    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[5])
    first.shapes.title.text = str(report.get("title", "Report"))
    deck = _Deck(prs, Inches, Pt)
    for si, sec in enumerate(report.get("sections", [])):
        sec_title = str(sec.get("title", ""))
        deck.new_slide(sec_title)
        for bi, block in enumerate(sec.get("blocks", [])):
            btype = block.get("type")
            if btype == "text":
                tb = deck.textbox(block["text"], 0.6, 0.7)
                tb.text_frame.word_wrap = True
            elif btype in ("params", "table"):
                header, rows = (params_rows if btype == "params" else table_rows)(block)
                nrows, ncols = len(rows) + 1, len(header)
                height = min(0.35 * nrows, 5.0)
                gt = deck.slide.shapes.add_table(
                    nrows, ncols, Inches(0.5), Inches(deck.top), Inches(9), Inches(height)
                ).table
                for j, h in enumerate(header):
                    gt.cell(0, j).text = h
                for i, row in enumerate(rows, start=1):
                    for j, c in enumerate(row):
                        gt.cell(i, j).text = str(c)
                deck.top += height + 0.3
            elif btype == "figure":
                plan = _plan(block, (si, bi), figures, warnings)
                deck.figure(block, plan, sec_title)
    buf = _io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
