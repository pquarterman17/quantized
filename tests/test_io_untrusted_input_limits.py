"""Hostile-input limits for the parsers (security audit, 2026-10-01).

Users open untrusted data files, so a parser must not let a small file claim
an unbounded amount of memory, read a file other than the one opened, or
smuggle markup into an exported document. Each test builds a minimal hostile
file and asserts the parser refuses it before doing the expensive work.

Where the real bound is large (hundreds of MB), the test patches the module's
limit down so the exploit stays small; ``raising=False`` keeps that patch
harmless on code that has no limit, so the test fails on the assertion.
"""

from __future__ import annotations

import re
import struct
import tracemalloc
import zipfile
from collections.abc import Callable, Iterator
from pathlib import Path
from typing import Any

import numpy as np
import pytest

from quantized.calc.report import ReportSheet, section, table_block, text_block
from quantized.datastruct import DataStruct
from quantized.io import _jcamp_asdf as asdf
from quantized.io import bruker_brml, excel, netcdf
from quantized.io.bruker_brml import import_bruker_brml
from quantized.io.excel import import_excel
from quantized.io.jcamp import import_jcamp
from quantized.io.netcdf import import_netcdf
from quantized.io.origin import format_origin_project_script, format_origin_script
from quantized.io.report_export import to_latex
from quantized.io.spc import import_spc


def _peak_bytes(fn: Callable[[], Any]) -> int:
    """Peak traced allocation while ``fn`` runs (numpy reports to tracemalloc)."""
    tracemalloc.start()
    try:
        try:
            fn()
        except ValueError:
            pass
        return tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()


@pytest.fixture
def h5py_mod() -> Iterator[Any]:
    yield pytest.importorskip("h5py")


# --------------------------------------------------------------------------
# NetCDF-4 / HDF5
# --------------------------------------------------------------------------
def test_netcdf4_refuses_external_storage(tmp_path: Path, h5py_mod: Any) -> None:
    """An HDF5 dataset can store its bytes in ANOTHER file by absolute path.

    Reading it would import that file's contents (any file the user can
    read, outside the allowed roots) as data."""
    secret = tmp_path / "secret.txt"
    secret.write_bytes(b"TOPSECRET-123456")
    path = tmp_path / "leak.nc"
    with h5py_mod.File(path, "w") as h:
        h.create_dataset("x", data=np.arange(16.0))
        h.create_dataset("y", shape=(16,), dtype="u1", external=[(str(secret), 0, 16)])
    with pytest.raises(ValueError, match="another file"):
        import_netcdf(path)


def test_netcdf4_refuses_virtual_dataset(tmp_path: Path, h5py_mod: Any) -> None:
    """A virtual dataset maps its data from other files the same way."""
    src = tmp_path / "src.h5"
    with h5py_mod.File(src, "w") as h:
        h.create_dataset("d", data=np.arange(16.0))
    layout = h5py_mod.VirtualLayout(shape=(16,), dtype="f8")
    layout[:] = h5py_mod.VirtualSource(str(src), "d", shape=(16,))
    path = tmp_path / "vds.nc"
    with h5py_mod.File(path, "w") as h:
        h.create_dataset("x", data=np.arange(16.0))
        h.create_virtual_dataset("y", layout)
    with pytest.raises(ValueError, match="another file"):
        import_netcdf(path)


def test_netcdf4_refuses_oversized_dataset(
    tmp_path: Path, h5py_mod: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A never-written dataset reads back as its fill value, so a ~1 KB file
    can declare billions of elements and make the read allocate them all."""
    monkeypatch.setattr(netcdf, "MAX_DATASET_ELEMENTS", 1000, raising=False)
    path = tmp_path / "bomb.nc"
    with h5py_mod.File(path, "w") as h:
        h.create_dataset("x", data=np.arange(16.0))
        h.create_dataset("y", shape=(5000,), dtype="f8", fillvalue=1.0)
    with pytest.raises(ValueError, match="elements"):
        import_netcdf(path)


@pytest.mark.parametrize("dtype", [np.dtype(("f8", (1000,))), np.dtype("S8000")])
def test_netcdf4_cap_counts_bytes_not_elements(
    tmp_path: Path, h5py_mod: Any, monkeypatch: pytest.MonkeyPatch, dtype: np.dtype[Any]
) -> None:
    """An element can be a whole array or a long fixed string, so a dataset
    under the element cap could still read back as terabytes of fill value."""
    monkeypatch.setattr(netcdf, "MAX_DATASET_ELEMENTS", 1000, raising=False)
    path = tmp_path / "wide.nc"
    with h5py_mod.File(path, "w") as h:
        h.create_dataset("x", data=np.arange(16.0))
        h.create_dataset("y", shape=(500,), dtype=dtype)  # 4 MB, never written
    with pytest.raises(ValueError, match="not read"):
        import_netcdf(path)


def _nc3_one_variable(dim_len: int) -> bytes:
    """A NetCDF-3 classic header declaring one float64 variable of ``dim_len``
    points, followed by only 64 bytes of data."""

    def name(s: bytes) -> bytes:
        return struct.pack(">i", len(s)) + s + b"\0" * (-len(s) % 4)

    out = b"CDF\x01" + struct.pack(">i", 0)
    out += struct.pack(">ii", 0x0A, 1) + name(b"t") + struct.pack(">i", dim_len)
    out += struct.pack(">ii", 0, 0)  # no global attributes
    out += struct.pack(">ii", 0x0B, 1) + name(b"v") + struct.pack(">ii", 1, 0)
    out += struct.pack(">ii", 0, 0)  # no variable attributes
    begin = len(out) + 12
    out += struct.pack(">iii", 6, (8 * dim_len) & 0x7FFFFFFF, begin)  # NC_DOUBLE
    return out + bytes(64)


def test_netcdf3_declared_size_is_not_allocated(tmp_path: Path) -> None:
    """The classic reader asked for each header-declared size in one read, and
    a buffered read allocates that much up front: 200 bytes -> 1.6 GB."""
    path = tmp_path / "bomb.nc"
    path.write_bytes(_nc3_one_variable(200_000_000))
    peak = _peak_bytes(lambda: import_netcdf(path))
    assert peak < 4 * 1024 * 1024, f"allocated {peak} bytes before refusing"
    with pytest.raises(ValueError):
        import_netcdf(path)


# --------------------------------------------------------------------------
# ZIP containers: Bruker .brml, Excel .xlsx
# --------------------------------------------------------------------------
def test_brml_refuses_decompression_bomb(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The scan XML is read whole; its uncompressed size is checked first."""
    monkeypatch.setattr(bruker_brml, "MAX_XML_BYTES", 1 << 20, raising=False)
    path = tmp_path / "bomb.brml"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("Experiment0/RawData0.xml", b"<RawData>" + b" " * (4 << 20) + b"</RawData>")
    assert path.stat().st_size < 64 * 1024  # ~1000:1 -- the bomb ratio
    with pytest.raises(ValueError, match="uncompressed"):
        import_bruker_brml(path)


def test_xlsx_refuses_sparse_cell_bomb(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """One far-away cell makes openpyxl pad every row and column before it:
    a few-KB workbook expands to rows x columns cells in memory."""
    openpyxl = pytest.importorskip("openpyxl")
    monkeypatch.setattr(excel, "MAX_CELLS", 100_000, raising=False)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"] = 1.0, 2.0
    ws.cell(row=1000, column=500, value=3.0)  # 500k cells once padded
    path = tmp_path / "sparse.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="cells"):
        import_excel(path)


def _xlsx_with_dimension(tmp_path: Path, rows: int, ref: str) -> Path:
    """A ``rows`` x 2 numeric sheet whose ``<dimension>`` tag says ``ref``."""
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["x", "y"])
    for i in range(rows):
        ws.append([float(i), float(2 * i)])
    plain = tmp_path / "plain.xlsx"
    wb.save(plain)
    path = tmp_path / "tagged.xlsx"
    with zipfile.ZipFile(plain) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                tag = f'<dimension ref="{ref}"'.encode()
                data, n = re.subn(rb'<dimension ref="[^"]*"', tag, data)
                assert n == 1
            dst.writestr(info.filename, data)
    return path


def test_xlsx_wrong_dimension_tag_does_not_count_against_the_cap(tmp_path: Path) -> None:
    """openpyxl pads every row to the tag's width, so ``A1:XFD...`` made a
    2-column sheet count 16,384 cells per row and refused it past ~2,048 rows."""
    path = _xlsx_with_dimension(tmp_path, 5000, "A1:XFD5001")
    ds = import_excel(path)
    assert ds.values.shape == (5000, 1)
    np.testing.assert_array_equal(ds.time, np.arange(5000.0))
    np.testing.assert_array_equal(ds.values[:, 0], 2 * np.arange(5000.0))


def test_xlsx_refuses_a_huge_real_grid(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    monkeypatch.setattr(excel, "MAX_CELLS", 100_000, raising=False)
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(400):
        ws.append([float(i + j) for j in range(300)])
    path = tmp_path / "grid.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="cells"):
        import_excel(path)


def test_xlsx_refuses_a_distant_row(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Missing rows before a far-down cell are yielded one by one; each counts."""
    openpyxl = pytest.importorskip("openpyxl")
    monkeypatch.setattr(excel, "MAX_CELLS", 100_000, raising=False)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"] = 1.0, 2.0
    ws["A200000"] = 3.0
    path = tmp_path / "tall.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="cells"):
        import_excel(path)


def test_xlsx_cap_counts_far_empty_styled_cells_as_scanned(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A styled empty cell far right makes openpyxl build the whole row; the
    trimmed row is small, but the work to read it still counts."""
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.styles import Font

    monkeypatch.setattr(excel, "MAX_CELLS", 100_000, raising=False)
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in range(1, 301):
        ws.cell(row=r, column=1, value=float(r))
        ws.cell(row=r, column=2, value=float(r))
        ws.cell(row=r, column=500).font = Font(bold=True)
    path = tmp_path / "styled.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="cells"):
        import_excel(path)


def test_xlsx_refuses_whole_parsed_part_bomb(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Non-worksheet parts (styles, shared strings) are parsed whole, so
    their uncompressed size is capped."""
    openpyxl = pytest.importorskip("openpyxl")
    monkeypatch.setattr(excel, "MAX_PART_BYTES", 1 << 20, raising=False)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["x", "label"])
    for i in range(5):
        ws.append([float(i), float(i * i)])
    plain = tmp_path / "plain.xlsx"
    wb.save(plain)
    path = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(plain) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "xl/styles.xml":
                data += b" " * (2 << 20)  # trailing whitespace: still valid XML
            dst.writestr(info.filename, data)
    with pytest.raises(ValueError, match="uncompressed"):
        import_excel(path)


_SST_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"


@pytest.mark.parametrize(
    ("member", "content_type"),
    [
        # A worksheet's own relationships part: read whole even in read-only mode.
        ("xl/worksheets/_rels/sheet1.xml.rels", None),
        # Shared strings are found through [Content_Types].xml, at any path.
        ("xl/worksheets/strings.xml", _SST_TYPE),
    ],
)
def test_xlsx_part_cap_is_not_bypassed_under_worksheets(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, member: str, content_type: str | None
) -> None:
    """Only the worksheets themselves are streamed; a part merely stored under
    ``xl/worksheets/`` is still parsed whole, so it is capped like any other."""
    openpyxl = pytest.importorskip("openpyxl")
    monkeypatch.setattr(excel, "MAX_PART_BYTES", 1 << 20, raising=False)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["x", "y"])
    for i in range(5):
        ws.append([float(i), float(i * i)])
    plain = tmp_path / "plain.xlsx"
    wb.save(plain)
    if content_type is None:
        root = "Relationships"
        ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    else:
        root = "sst"
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    bomb = f'<{root} xmlns="{ns}">'.encode() + b" " * (2 << 20) + f"</{root}>".encode()
    path = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(plain) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "[Content_Types].xml" and content_type is not None:
                override = f'<Override PartName="/{member}" ContentType="{content_type}"/>'
                data = data.replace(b"</Types>", override.encode() + b"</Types>")
            dst.writestr(info.filename, data)
        dst.writestr(member, bomb)
    with pytest.raises(ValueError, match="uncompressed"):
        import_excel(path)


# --------------------------------------------------------------------------
# Header-declared counts: SPC fnpts, JCAMP DUP run-length
# --------------------------------------------------------------------------
def _spc_even(fnpts: int, n_data_bytes: int) -> bytes:
    """A modern (0x4B) evenly-spaced single-subfile SPC header + payload."""
    fmt = "<BBBbIddIBBBBI9s9sh32s130s30sIIBBhf48sfIfB187s"
    fields = list(struct.unpack(fmt, bytes(struct.calcsize(fmt))))  # all zero
    # ftflgs, fversn, fexper, fexp, fnpts, ffirst, flast, fnsub
    fields[:8] = [0, 0x4B, 0, 0, fnpts, 0.0, 1.0, 1]
    return struct.pack(fmt, *fields) + bytes(32) + bytes(n_data_bytes)


def test_spc_refuses_point_count_the_file_cannot_hold(tmp_path: Path) -> None:
    """fnpts is a header uint32: the even-x axis was allocated from it before
    the data length was checked (4e9 points -> ~32 GB from a 1 KB file)."""
    path = tmp_path / "bomb.spc"
    path.write_bytes(_spc_even(5_000_000, 64))
    peak = _peak_bytes(lambda: import_spc(path))
    assert peak < 4 * 1024 * 1024, f"allocated {peak} bytes before refusing"
    with pytest.raises(ValueError, match="points"):
        import_spc(path)


def test_jcamp_dup_run_cannot_exceed_declared_npoints(tmp_path: Path) -> None:
    """A DUP token is a run length: ``s9999999`` alone is ~100M ordinates."""
    path = tmp_path / "bomb.jdx"
    path.write_text(
        "##TITLE=bomb\n##JCAMP-DX=4.24\n##DATA TYPE=INFRARED SPECTRUM\n"
        "##XUNITS=1/CM\n##YUNITS=ABSORBANCE\n##FIRSTX=0\n##LASTX=2\n"
        "##XFACTOR=1\n##YFACTOR=1\n##NPOINTS=3\n##FIRSTY=1\n"
        "##XYDATA=(X++(Y..Y))\n0 1s999999\n##END=\n",
        encoding="latin-1",
    )
    peak = _peak_bytes(lambda: import_jcamp(path))
    assert peak < 4 * 1024 * 1024, f"allocated {peak} bytes before refusing"
    with pytest.raises(ValueError, match="DUP run"):
        import_jcamp(path)


def test_asdf_dup_run_is_capped_without_npoints(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(asdf, "MAX_ORDINATES", 1000, raising=False)
    with pytest.raises(ValueError, match="DUP run"):
        asdf.decode_xydata(["0 1s99999"])


# --------------------------------------------------------------------------
# Exported documents: LaTeX report, LabTalk .ogs script
# --------------------------------------------------------------------------
def test_latex_escapes_backslash_commands() -> None:
    r"""File-derived text (a column label, a note) must not reach LaTeX as a
    command: ``\input /etc/passwd`` would inline that file into the PDF."""
    hostile = r"\input /etc/passwd \immediate\write18"
    rep = ReportSheet(title=hostile, sections=(
        section(hostile, [text_block(hostile), table_block([hostile, "b"], [[hostile, 1]])]),
    )).to_dict()
    tex = to_latex(rep)
    assert r"\input" not in tex
    assert r"\immediate" not in tex
    assert r"\textbackslash{}input" in tex


def _ds(label: str, unit: str = "", **meta: Any) -> DataStruct:
    return DataStruct(
        time=np.array([0.0, 1.0]),
        values=np.array([[1.0], [2.0]]),
        labels=(label,),
        units=(unit,),
        metadata=meta,
    )


def _statements(ogs: str) -> list[str]:
    return [ln.strip() for ln in ogs.splitlines()]


def test_labtalk_script_cannot_gain_lines_from_a_label() -> None:
    """A newline in a label or unit ended the string literal's line, so the
    rest ran as its own LabTalk statement (``run -e`` starts a program)."""
    payload = 'x\nrun -e calc.exe;\r\ntype "pwned";'
    _csv, ogs = format_origin_script(
        _ds(payload, payload, x_column_name=payload, x_column_unit=payload),
        make_graph=True,
    )
    assert not any(s.startswith(("run ", "type ")) for s in _statements(ogs))
    _csvs, ogs2 = format_origin_project_script(
        [(_ds(payload, origin_book_long=payload), "book")]
    )
    assert not any(s.startswith(("run ", "type ")) for s in _statements(ogs2))


def test_labtalk_string_literal_holds_no_inner_double_quote() -> None:
    """LabTalk has no backslash escape (live-verified 2026-07-04,
    docs/origin_re/validation_log.md): ``\\"`` does not keep a quote inside
    the literal. A label's quote becomes ``'``, as ``origin_com`` does."""
    payload = 'a"; run -e calc.exe; "b'
    _csv, ogs = format_origin_script(
        _ds(payload, payload, x_column_name=payload, x_column_unit=payload),
        make_graph=True,
    )
    literals = [ln for ln in ogs.splitlines() if "a" in ln and "calc" in ln]
    assert literals, "the label never reached the script"
    for line in literals:
        assert line.count('"') == 2, line  # just the literal's own pair
        assert "a'; run -e calc.exe; 'b" in line
